from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

output_path = Path(__file__).resolve().parent / "STRATEGY_SUMMARY.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = 720

p = doc.add_paragraph()
r = p.add_run("SB-PR-DQP Strategy & Workflow Summary")
r.bold = True
r.font.size = Pt(20)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

items = [
    (
        "1. Overview",
        "This repo is built for intraday trading on Zerodha Kite. The main focus is intraday futures, with an additional intraday options mode. The bot ingests live market data, aggregates ticks into 1-minute bars, applies technical indicators, manages risk, logs every important decision, and exits all open positions before market close.",
    ),
    (
        "2. Futures strategy",
        "The futures strategy is a trend-following system built around EMA crossover and RSI confirmation. In strategy.py, the bot checks whether the fast EMA is above/below the slow EMA, whether price is above/below the 20-period SMA, and whether RSI is in a valid trend zone. A BUY is favored when price is trending upward and RSI supports momentum; a SELL is favored when price is trending lower and RSI confirms bearish momentum. It also allows breakout-style entries when the recent high/low is broken. This keeps the system focused on high-quality trend entries rather than random choppy moves.",
    ),
    (
        "3. Options strategy",
        "The options mode follows the same intraday workflow but trades option legs instead of futures. The preferred option leg is mapped by signal direction: BUY prefers CE and SELL prefers PE. Before an option trade is allowed, the bot checks premium, volume, OI, and IV quality to avoid weak contracts. It only allows the expected option side for each underlying and rejects low-liquidity or bad-quality option contracts. In intraday_options mode, the configured stop and target are applied to the option premium itself, not just the underlying spot. The bot also supports intraday_both to run futures and options together.",
    ),
    (
        "4. Risk management",
        "Risk control is handled by IntradayManager. It calculates risk based on account size and configured risk per trade, tracks a daily loss cap, applies loss-streak and symbol-level safeguards, and monitors active positions. Futures position sizing is derived from the stop-loss distance, contract multiplier, and lot size. It sets SL and TP for each trade, auto-closes on TP/SL, and force-closes all open positions at 3:15 PM IST to avoid overnight exposure.",
    ),
    (
        "5. Repo workflow / execution steps",
        "1. Set KITE_API_KEY and KITE_ACCESS_TOKEN as environment variables.\n2. Copy the sample config and set trading_mode to intraday_futures, intraday_options, or intraday_both.\n3. Configure the symbol universe using futures_underlyings and/or options_underlyings.\n4. Run the bot with python run_kite_bot.py.\n5. The bot warms up recent historical candles and subscribes to live Kite ticks.\n6. Tick data is aggregated into 1-minute bars.\n7. Every bar is evaluated using EMA + RSI + trend logic and optional market/news filters.\n8. Valid signals generate entries with stop-loss and take-profit levels.\n9. The bot monitors real-time P&L and exits on stop-loss, target, or forced market-close.\n10. Details are written to trades.jsonl and decision_log.jsonl for audit and review.\n11. Every Friday, run python analyze_weekly.py to measure weekly performance by symbol, hour, and trade outcome.\n12. Fill weekly_input_template.json based on the report and tune config/strategy for the next cycle.",
    ),
    (
        "6. Weekly improvement loop",
        "This repo follows a disciplined weekly improvement loop. The trader runs the bot daily, reviews the report on Friday, fills in the feedback form, and then updates configuration or strategy logic for the following week. The goal is to improve performance through actual evidence rather than repeated guesswork.",
    ),
    (
        "7. Final summary",
        "In short, this repo is using a disciplined intraday trend-following system: EMA + RSI for futures, filtered option selection for options, strong risk controls, accuracy filters, and a forced 3:15 PM exit. The operational model is: monitor live bars, take only clean setups, protect capital, log every decision, and refine every week using real market results.",
    ),
]

for heading, body in items:
    doc.add_heading(heading, level=1)
    for line in body.split("\n"):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.save(output_path)
print(f"Created: {output_path}")
print(f"Exists: {output_path.exists()}")
